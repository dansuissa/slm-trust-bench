"""Local adapters for frozen author data. Preparation performs no model inference."""
import ast
from collections import Counter, defaultdict
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from .io import digest, validate_items, write_json, write_jsonl


def key(text):
    return sha256(" ".join(text.split()).encode()).hexdigest()


def item(identifier, group, dataset, task, split, messages, references=None, should_abstain=None, **extra):
    return dict(item_id=identifier, group_id=group, dataset=dataset, task=task, split=split,
                messages=messages, references=references or [], should_abstain=should_abstain, **extra)


def vrr_prompt(source, values):
    """Read prompt literals from archived author code without importing or executing it."""
    tree = ast.parse(source.read_text(encoding="utf-8"))
    definitions = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "value_definitions" for t in node.targets):
            definitions = ast.literal_eval(node.value)
    values = dict(values)
    if definitions:
        # Preserve the author's explicit fallback and flag it in the preparation audit.
        values.update(value1_definition=definitions.get(values["value1"], 'Definition not found.'),
                      value2_definition=definitions.get(values["value2"], 'Definition not found.'))
    prompt_node = next(node.value for node in ast.walk(tree) if isinstance(node, ast.Assign)
                       and any(isinstance(t, ast.Name) and t.id == "prompt" for t in node.targets))
    pieces = []
    for part in prompt_node.values:
        if isinstance(part, ast.Constant):
            pieces.append(part.value)
        elif isinstance(part, ast.FormattedValue) and isinstance(part.value, ast.Name):
            pieces.append(str(values[part.value.id]))
        else:
            raise ValueError("Unsupported upstream prompt expression")
    system = None
    for node in ast.walk(tree):
        if isinstance(node, ast.Dict):
            try:
                literal = ast.literal_eval(node)
            except (ValueError, TypeError):
                continue
            if literal.get("role") == "system":
                system = literal["content"]
    if system is None:
        raise ValueError("Upstream system prompt missing")
    return [{"role": "system", "content": system}, {"role": "user", "content": "".join(pieces)}]


def prepare(root):
    raw = root / "data/raw"
    out = root / "data/prepared"
    manifest = json.loads((root / "manifest.json").read_text(encoding='utf-8'))
    for source in manifest.get('prompt_sources', []):
        path = root / 'data/sources' / source['repository'].replace('/', '__') / 'files' / source['file']
        if digest(path) != source['sha256']:
            raise ValueError(f"Prompt source hash mismatch: {source['file']}")
    filenames = {"facttest_training_data": "facttest_training_data.json",
                 "facttest_ID_test_pararel": "facttest_ID_test_pararel.json",
                 "facttest_OOD_test_pararel": "facttest_OOD_test_pararel.json",
                 "value_reasoning": "value_reasoning.zip", "squad_v2_validation": "squad_v2_validation.parquet"}
    for entry in manifest["datasets"]:
        if digest(raw / filenames[entry["id"]]) != entry["sha256"]:
            raise ValueError(f"Raw data hash mismatch: {entry['id']}")
    summary = {"raw_hashes_verified": True, "datasets": {}}
    factual_sets = {}
    for name, split in [("training_data", "calibration"), ("ID_test_pararel", "test"), ("OOD_test_pararel", "shift_test")]:
        rows = json.loads((raw / f"facttest_{name}.json").read_text(encoding='utf-8'))
        if any(not isinstance(row, list) or len(row) != 3 for row in rows):
            raise ValueError('Unexpected ParaRel record format')
        items = [item(f"pararel:{name}:{i}", key(question), "pararel", "factual", split,
                      [{"role": "user", "content": question}], [answer], relation=relation)
                 for i, (question, answer, relation) in enumerate(rows)]
        validate_items(items)
        path = out / f"pararel_{split}.jsonl"
        write_jsonl(path, items)
        factual_sets[split] = items
        summary["datasets"][path.name] = {"items": len(items), "sha256": digest(path),
                                           "duplicate_question_groups": len(items) - len({x['group_id'] for x in items})}
    summary["pararel_question_group_overlap"] = {
        f"{a}__{b}": len({x['group_id'] for x in factual_sets[a]} & {x['group_id'] for x in factual_sets[b]})
        for a, b in [("calibration", "test"), ("calibration", "shift_test"), ("test", "shift_test")]}
    clean_sets = {}
    all_calibration_groups = {x['group_id'] for x in factual_sets['calibration']}
    for split in ['calibration', 'test']:
        groups = defaultdict(list)
        for row in factual_sets[split]:
            groups[row['group_id']].append(row)
        clean = []
        ambiguous = 0
        shared = 0
        for group, rows in sorted(groups.items()):
            if split == 'test' and group in all_calibration_groups:
                shared += 1
                continue
            if len({tuple(row['references']) for row in rows}) > 1:
                ambiguous += 1
                continue
            row = dict(rows[0], split=split + '_clean', source_multiplicity=len(rows))
            clean.append(row)
        path = out / f'pararel_{split}_clean.jsonl'
        write_jsonl(path, clean)
        clean_sets[split] = clean
        summary['datasets'][path.name] = {'items': len(clean), 'sha256': digest(path),
                                         'shared_question_groups_removed': shared,
                                         'multiple_reference_groups_excluded_pending_review': ambiguous}

    import pyarrow.parquet as parquet
    rows = parquet.read_table(raw / "squad_v2_validation.parquet").to_pylist()
    items = []
    for row in rows:
        references = row["answers"]["text"]
        prompt = f"Respond to the question using only information given in the context.\nContext: {row['context']}\nQuestion: {row['question']}"
        items.append(item(row['id'], key(row['context']), "squad2", "abstention", "published_validation",
                          [{"role": "user", "content": prompt}], references, not bool(references)))
    validate_items(items)
    # Group by passage before forming research development/calibration/test splits.
    grouped = {group: int(sha256(("split-seed-679:" + group).encode()).hexdigest(), 16) % 10
               for group in {x['group_id'] for x in items}}
    for row in items:
        bucket = grouped[row['group_id']]
        row['research_split'] = 'development' if bucket < 2 else 'calibration' if bucket < 5 else 'test'
    path = out / "squad2.jsonl"
    write_jsonl(path, items)
    summary["datasets"][path.name] = {"items": len(items), "sha256": digest(path),
        "answerability_counts": dict(Counter(str(x['should_abstain']) for x in items)),
        "research_split_counts": dict(Counter(x['research_split'] for x in items)),
        "independent_passage_groups": len(grouped)}

    from openpyxl import load_workbook
    upstream = root / 'data/sources/Giovanna-SH__Value_Reasoning_Reliability/files/src/models'
    with ZipFile(raw / "value_reasoning.zip") as archive:
        for kind in ["PVQ", "INVP"]:
            items = []
            for name in sorted(archive.namelist()):
                if not name.endswith('.xlsx'):
                    continue
                if kind == 'PVQ' and not name.endswith('/PVQ40/PVQ40.xlsx'):
                    continue
                if kind == 'INVP' and '/INVP/original data/' not in name:
                    continue
                workbook = load_workbook(BytesIO(archive.read(name)), read_only=True, data_only=True)
                iterator = workbook.active.iter_rows(values_only=True)
                headers = next(iterator)
                source = upstream / ('en_originQfile.py' if kind == 'PVQ' else 'en_defination-putconfident.py')
                for i, cells in enumerate(iterator):
                    values = dict(zip(headers, cells))
                    if any(v is None for v in values.values()):
                        raise ValueError(f"Missing workbook cell: {name}:{i + 2}")
                    params = dict(context=values['Q'], option=values['A']) if kind == 'PVQ' else dict(
                        context=values['scene'], value1=values['value1'], value2=values['value2'], d1=values['D1'], d2=values['D2'])
                    options = {str(n): str(n) for n in range(1, 7)} if kind == 'PVQ' else {'A': 'action_1', 'B': 'action_2'}
                    identifier = f'{kind}:{Path(name).stem}:{i}'
                    messages = vrr_prompt(source, params)
                    items.append(item(identifier, identifier, kind.lower(), 'value', 'original',
                                      messages, option_map=options, source_file=name, source_row=i + 2,
                                      prompt_source_sha256=digest(source),
                                      missing_definition='Definition not found.' in messages[-1]['content']))
                workbook.close()
            validate_items(items)
            path = out / f'{kind.lower()}.jsonl'
            write_jsonl(path, items)
            summary['datasets'][path.name] = {'items': len(items), 'sha256': digest(path),
                                             'prompt_protocol': 'author release literal reconstruction',
                                             'items_with_missing_value_definition': sum(x['missing_definition'] for x in items)}
    write_json(out / 'audit.json', summary)
    return summary
