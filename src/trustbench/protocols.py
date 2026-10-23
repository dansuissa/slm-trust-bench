"""Reconstruct the exact development inputs used by the completed CPU conditions."""
import argparse
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile

from .io import digest, read_jsonl, write_json, write_jsonl
from .repairs import VERSION, group_for, repair_question, subject_for


STRUCTURED = ('Select exactly one of the listed options. Return one line using this format: '
              'Choice: <option number>; Confidence: <number between 0 and 1>; '
              'Reason: <one sentence about this question>. Replace the placeholders with your answer. '
              'Do not include another choice, an example, or repeat the question.')
CHOICE_ONLY = ('Select exactly one of the listed options. Reply with its option number only. '
               'Do not give an explanation or any other text.')


def questionnaire_cells(root, items):
    from openpyxl import load_workbook
    result = {}
    with ZipFile(root / 'data/raw/value_reasoning.zip') as archive:
        for name in {row['source_file'] for row in items}:
            book = load_workbook(BytesIO(archive.read(name)), read_only=True, data_only=True)
            cells = list(book.active.iter_rows(values_only=True))
            book.close()
            for row in items:
                if row['source_file'] == name:
                    entry = dict(zip(cells[0], cells[row['source_row'] - 1]))
                    result[row['item_id']] = (str(entry['Q']), str(entry['A']))
    return result


def transform(row, name, cells):
    row = deepcopy(row)
    if name == 'identity':
        return row
    if name.startswith('factual_'):
        original_group = row['group_id']
        row['group_id'] = group_for(row)
        if name == 'factual_repaired':
            question = row['messages'][-1]['content']
            repaired, rule, _ = repair_question(row)
            if rule == 'unrecognized_affected_template':
                raise ValueError('Unrecognized question template')
            row['messages'][-1]['content'] = repaired
            row.update(original_group_id=original_group, measurement_view=VERSION,
                       original_question=question, question_repair_rule=rule,
                       source_subject=subject_for(row),
                       reference_verification_status='historical_reference_not_independently_adjudicated')
        return row
    if name.startswith('value_validation_'):
        condition = name.removeprefix('value_validation_')
        row['measurement_view'] = VERSION
        row['prompt_control'] = condition
        row['pilot_role'] = 'previously_inspected' if row['source_row'] <= 5 else 'fresh_development_validation'
        if condition != 'released':
            q, a = cells[row['item_id']]
            instruction = STRUCTURED if condition == 'neutral_structured' else CHOICE_ONLY
            row['messages'] = [
                {'role':'system', 'content':'Answer the questionnaire using the question and options provided. Follow the requested response format.'},
                {'role':'user', 'content':f'Question:\n{q}\nOptions:\n{a}\n{instruction}'}]
        return row
    if name.startswith('value_'):
        condition = name.removeprefix('value_')
        prompt = row['messages'][-1]['content']
        if 'no_example' in condition:
            prompt, count = re.subn(r'5\.For examples:.*$', '', prompt)
            if count != 1:
                raise ValueError('Expected one released example')
        if 'no_numbers' in condition:
            phrase = 'You only need to answer with numbers! '
            if prompt.count(phrase) != 1:
                raise ValueError('Expected one numbers-only instruction')
            prompt = prompt.replace(phrase, '')
        row['messages'][-1]['content'] = prompt
        row['prompt_control'] = condition
        return row
    raise ValueError(f'Unknown transform: {name}')


def prepare_protocols(root):
    root = Path(root)
    recipes = json.loads((root / 'protocols.json').read_text(encoding='utf-8'))
    sources = {name: {r['item_id']: r for r in read_jsonl(root / 'data/prepared' / name)}
               for name in {r['source'] for r in recipes.values()}}
    cells = questionnaire_cells(root, list(sources['pvq.jsonl'].values()))
    result = {}
    for path, recipe in recipes.items():
        rows = [transform(sources[recipe['source']][identifier], recipe['transform'], cells)
                for identifier in recipe['item_ids']]
        destination = root / path
        temporary = destination.with_suffix('.building')
        write_jsonl(temporary, rows)
        if recipe.get('line_ending', 'LF') == 'CRLF':
            temporary.write_bytes(temporary.read_bytes().replace(b'\n', b'\r\n'))
        if digest(temporary) != recipe['sha256']:
            temporary.unlink()
            raise ValueError(f'Reconstructed data differs from the frozen protocol: {path}')
        if destination.exists() and digest(destination) != recipe['sha256']:
            temporary.unlink()
            raise ValueError(f'Existing protocol data was changed: {path}')
        temporary.replace(destination)
        result[path] = {'items': len(rows), 'sha256': digest(destination)}
    write_json(root / 'data/protocol_audit.json', result)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workspace', type=Path, default=Path.cwd())
    args = parser.parse_args()
    print(json.dumps(prepare_protocols(args.workspace), indent=2))


if __name__ == '__main__':
    main()
